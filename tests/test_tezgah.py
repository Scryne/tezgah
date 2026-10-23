import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tezgah import kontrol
from tezgah.donustur import donustur
from tezgah.sayi import coz, sayi_mi

VERI = Path(__file__).parent / "veri"


def dogru(ad):
    return json.loads((VERI / f"{ad}.json").read_text(encoding="utf-8"))


def tablo(satirlar, kaynak="x", sayfa=1):
    return {"kaynak": kaynak, "sayfa": sayfa, "satirlar": satirlar}


def siparis(tmp_path, *dosyalar):
    k = tmp_path / "siparis"
    (k / "girdi").mkdir(parents=True)
    for d in dosyalar:
        shutil.copy(VERI / d, k / "girdi" / d)
    return k


@pytest.mark.parametrize("metin,beklenen", [
    ("1.234,56", 1234.56), ("191,08", 191.08), ("25", 25), ("-5,5", -5.5), ("0,50", 0.5), ("2.920.163,54", 2920163.54),
    ("191.08", None), ("0110", None), ("14.10.2026", None), ("12'li", None), ("UR-1000", None), ("", None),
])
def test_sayi(metin, beklenen):
    assert coz(metin) == beklenen
    assert sayi_mi(metin) == (beklenen is not None)


@pytest.mark.parametrize("ad", ["a_fiyat_listesi", "b_fatura", "c_uzun_iki_sayfa"])
def test_dijital_pdf_excele_birebir(tmp_path, ad):
    k = siparis(tmp_path, f"{ad}.pdf")
    s = donustur(k)
    assert s.durum == "TEMIZ", s.bulgular
    ws = load_workbook(s.ciktilar[0]).active
    g = dogru(ad)
    satirlar = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    assert len(ws.sheetnames if hasattr(ws, "sheetnames") else [ws]) == 1
    assert [str(c) for c in satirlar[0]] == g["basliklar"]  # iki sayfalık tablo tek sayfaya, başlık bir kez
    assert len(satirlar) == len(g["satirlar"]) + 1
    for beklenen, gelen in zip(g["satirlar"], satirlar[1:]):
        for b, x in zip(beklenen, gelen):
            assert (coz(b) if coz(b) is not None else (b or None)) == x


def test_tarama_claude_okumasini_bekler_sonra_teslim_eder(tmp_path):
    k = siparis(tmp_path, "a_fiyat_listesi_tarama_s1.jpg")
    s = donustur(k, ikinci_okuma=False)
    assert s.durum == "BEKLIYOR" and len(s.bekleyen) == 1
    satirlar = json.loads((VERI / "claude_a.json").read_text(encoding="utf-8"))
    (k / "okuma_claude.json").write_text(json.dumps(
        {"tablolar": [tablo(satirlar, "a_fiyat_listesi_tarama_s1")]}, ensure_ascii=False), encoding="utf-8")
    s = donustur(k, ikinci_okuma=False, word=True)
    assert s.durum == "TEMIZ", s.bulgular
    assert all(p.exists() for p in s.ciktilar) and len(s.ciktilar) == 2


def test_ikinci_okuma_uyusmazligi_teslimi_durdurur_dogrulaninca_gecer(tmp_path):
    k = siparis(tmp_path, "a_fiyat_listesi_tarama_s1.jpg")
    satirlar = json.loads((VERI / "claude_a.json").read_text(encoding="utf-8"))
    ikinci = [list(r) for r in satirlar]
    ikinci[3][1] = "Kağıt havlu (24'lü)"  # metindeki rakam farklı: hesap bunu doğrulayamaz
    kaynak = "a_fiyat_listesi_tarama_s1"
    (k / "okuma_docling.json").write_text(json.dumps({"tablolar": [tablo(ikinci, kaynak)]}), encoding="utf-8")
    claude = {"tablolar": [tablo(satirlar, kaynak)]}
    (k / "okuma_claude.json").write_text(json.dumps(claude), encoding="utf-8")
    s = donustur(k)
    assert s.durum == "INCELE"
    assert [(b.tur, b.satir, b.sutun, b.deger) for b in s.bulgular] == [("uyusmazlik", 3, 1, "Kağıt havlu (12'li)")]
    assert not (k / "cikti").exists()
    claude["dogrulandi"] = [{"kaynak": kaynak, "sayfa": 1, "satir": 3, "sutun": 1}]
    (k / "okuma_claude.json").write_text(json.dumps(claude), encoding="utf-8")
    assert donustur(k).durum == "TEMIZ"


def fatura():
    g = dogru("b_fatura")
    return tablo([g["basliklar"]] + [list(r) for r in g["satirlar"]])


def test_aritmetik_temiz_faturada_sessiz():
    assert kontrol.aritmetik(fatura()) == [] and kontrol.bicim(fatura()) == []


def test_aritmetik_bozuk_tutari_yakalar():
    t = fatura()
    t["satirlar"][5][5] = "356.449,82"
    assert [(b.tur, b.satir, b.sutun) for b in kontrol.aritmetik(t)] == [("carpim", 5, 5), ("toplam", 13, 5)]


def test_aritmetik_bozuk_genel_toplami_yakalar():
    t = fatura()
    t["satirlar"][-1][5] = "3.504.196,52"
    assert [(b.tur, b.satir) for b in kontrol.aritmetik(t)] == [("toplam", 15)]


def test_ondalik_nokta_okumasi_sessizce_metne_donmez():
    t = fatura()
    t["satirlar"][2][4] = "191.08"
    assert [(b.tur, b.satir, b.sutun) for b in kontrol.bicim(t)] == [("bicim", 2, 4)]


def test_karsilastirma_harf_farkini_affeder_rakami_affetmez():
    claude = [tablo([["UR-1007", "Şerit metre 5 m", "adet", "108", "191,08"]])]
    docling = [tablo([["UR-1007", "Serit metre 5m", "adet", "108", "191.08"]])]
    assert [(b.sutun, b.deger) for b in kontrol.karsilastir(claude, docling)] == [(4, "191,08")]


def test_metin_gurultusu_affedilir_metindeki_rakam_affedilmez():
    claude = [tablo([["UR-1154", "Özel cıvata M10", "adet"], ["UR-1063", "Üçgen reflektör", "adet"],
                     ["UR-1140", "Ğ testi numune kabı", "paket"], ["UR-1155", "Özel cıvata M10", "adet"]])]
    docling = [tablo([["UR-1154", "Ozel cıvata M1O", "adet"], ["UR-1063", "reflektör Üçgen", "adet"],
                      ["UR-1140", "testi numune kabı", "paket"], ["UR-1155", "Özel cıvata M16", "adet"]])]
    assert [(b.satir, b.deger) for b in kontrol.karsilastir(claude, docling)] == [(3, "Özel cıvata M10")]


def test_hesabi_tutan_satirda_rakam_uyusmazligi_alarm_degil_tutmayanda_alarm():
    t = fatura()
    t["satirlar"] = t["satirlar"][:6]
    ikinci = [list(r) for r in t["satirlar"]]
    ikinci[2][5] = ""          # docling tutarı kaçırdı; Claude'un satırında hesap tutuyor
    ikinci[3][4] = "2.680,90"  # docling farklı okudu; hesap tutuyor
    assert kontrol.karsilastir([t], [tablo(ikinci)]) == []
    t["satirlar"][4][4] = "3.338,38"  # Claude yanlış okudu: hesap tutmaz, ikinci okuma da farklı
    assert [(b.satir, b.sutun) for b in kontrol.karsilastir([t], [tablo(ikinci)])] == [(4, 4)]
