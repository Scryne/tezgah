import json

from openpyxl import Workbook, load_workbook

from tezgah import trendyol

ACIKLAMA = (
    "El yapımı inci kolye, deniz kabuğu charmı ve ince zinciriyle günlük kullanıma uygun. Zincir paslanmaz çelik, "
    "uzunluğu 45 cm. Suyla ve parfümle temastan kaçının, kuru bezle silin. Hediye kutusunda gönderilir."
)


def siparis(tmp_path):
    (tmp_path / "girdi").mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["Trendyol ürün listesi"])  # başlık satırı ilk satır olmayabilir
    ws.append(["Barkod", "Model Kodu", "Ürün Adı", "Ürün Açıklaması", "Renk", "Materyal", "Fiyat"])
    ws.append(["868001", "K-1", "İNCİLİ KOLYE", None, "Beyaz", None, 399])
    ws.append(["868002", "K-2", "deniz kabuğu kolye paslanmaz çelik", "kolye", None, None, 299])
    wb.save(tmp_path / "girdi" / "urunler.xlsx")
    trendyol.ac(tmp_path)
    return tmp_path


def duzenle(klasor, **degisiklik):
    d = {
        "urunler": [
            {"barkod": "868001", "baslik": "Kadın El Yapımı İnci Kolye Deniz Kabuğu Charmlı İnce Zincir", "aciklama": ACIKLAMA,
             "ozellikler": {"Renk": {"deger": "Beyaz", "kaynak": "urun"}, "Materyal": {"deger": "Paslanmaz Çelik", "kaynak": "satici"}}},
            {"barkod": "868002", "baslik": "Kadın Deniz Kabuğu Kolye Paslanmaz Çelik Zincir", "aciklama": ACIKLAMA,
             "ozellikler": {"Materyal": {"deger": "Paslanmaz Çelik", "kaynak": "urun"}}},
        ],
        "sorular": [],
    }
    d.update(degisiklik)
    (klasor / trendyol.DUZENLEME).write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    (klasor / trendyol.SATICI).write_text(json.dumps({"868001": {"Materyal": "Paslanmaz Çelik"}}, ensure_ascii=False), encoding="utf-8")
    return d


def test_ac_basligi_ikinci_satirda_bulur_ve_sablon_yazar(tmp_path):
    k = siparis(tmp_path)
    urunler = json.loads((k / trendyol.URUNLER).read_text(encoding="utf-8"))
    assert [u["barkod"] for u in urunler] == ["868001", "868002"]
    assert urunler[0]["alanlar"]["Ürün Adı"] == "İNCİLİ KOLYE"
    sablon = json.loads((k / trendyol.DUZENLEME).read_text(encoding="utf-8"))
    assert [u["barkod"] for u in sablon["urunler"]] == ["868001", "868002"]


def test_temiz_duzenleme_excel_ve_metin_teslim_eder(tmp_path):
    k = siparis(tmp_path)
    duzenle(k)
    s = trendyol.teslim(k)
    assert s.durum == "TEMIZ", s.bulgular
    wb = load_workbook(k / "teslim" / "urunler-duzenlendi.xlsx")
    ws = wb.active
    assert ws["C3"].value == "Kadın El Yapımı İnci Kolye Deniz Kabuğu Charmlı İnce Zincir"
    assert ws["C3"].fill.start_color.rgb == "FFE8F3D6"
    assert ws["F3"].value == "Paslanmaz Çelik"
    assert ws["G3"].value == 399  # dokunulmayan sütun aynı
    assert ws["E3"].fill.start_color.rgb != "FFE8F3D6"  # Renk değişmedi, boyanmadı
    degisen = list(wb["Değişiklikler"].iter_rows(values_only=True))
    assert ("868001", "Ürün Adı", "İNCİLİ KOLYE", "Kadın El Yapımı İnci Kolye Deniz Kabuğu Charmlı İnce Zincir") in degisen
    assert "Barkod: 868002" in (k / "teslim" / "urunler.md").read_text(encoding="utf-8")


def test_kusurlu_duzenleme_teslim_edilmez(tmp_path):
    k = siparis(tmp_path)
    d = duzenle(k)
    u1, u2 = d["urunler"]
    u1["baslik"] = "KADIN EL YAPIMI İNCİ KOLYE DENİZ KABUĞU"
    u1["aciklama"] = ACIKLAMA + " Sipariş için 0532 123 45 67 numarasını arayın."
    u1["ozellikler"]["Taş Cinsi"] = {"deger": "Sedef", "kaynak": "urun"}  # ürün verisinde yok: uydurma
    u1["ozellikler"]["Gramaj"] = {"deger": "8 gr", "kaynak": "gorsel"}  # görselden okunamaz
    u2["aciklama"] = "kısa"
    d["urunler"] = [u1, {"barkod": "999999", "baslik": "x" * 20, "aciklama": ACIKLAMA, "ozellikler": {}}]
    (k / trendyol.DUZENLEME).write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    s = trendyol.teslim(k)
    assert s.durum == "INCELE"
    metin = "\n".join(s.bulgular)
    for beklenen in ("büyük harf", "iletişim", "'Sedef' ürün verisinde geçmiyor", "görselden okunamaz",
                     "999999: girdide olmayan", "868002: düzenlemede yok"):
        assert beklenen in metin, beklenen
    assert not (k / "teslim").exists()


def test_satici_cevabi_yoksa_kaynak_satici_gecmez(tmp_path):
    k = siparis(tmp_path)
    duzenle(k)
    (k / trendyol.SATICI).write_text("{}", encoding="utf-8")
    s = trendyol.kontrol(k)
    assert any("satıcı cevaplarında yok" in b for b in s.bulgular)


def test_soru_varsa_bekler_ve_mesaj_uretir(tmp_path):
    k = siparis(tmp_path)
    duzenle(k, sorular=["868002 kolyenin zincir uzunluğu kaç cm?"])
    s = trendyol.teslim(k)
    assert s.durum == "BEKLIYOR"
    assert "1. 868002 kolyenin zincir uzunluğu" in trendyol.soru_mesaji(s.sorular)
    assert not (k / "teslim").exists()
