"""Portfolyo görselleri (1280x720): gerçek tarama + `tezgah donustur` çıktısı .xlsx'ten çizilen tablo."""
import base64
import io
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from playwright.sync_api import sync_playwright

KOK = Path(__file__).resolve().parents[3]
OUT = Path(__file__).parent
CHROME = Path.home() / "AppData/Local/ms-playwright/chromium-1243/chrome-win64/chrome.exe"


def veri_uri(yol, kirp=None, genislik=None):
    im = Image.open(yol).convert("RGB")
    if kirp:
        im = im.crop(kirp)
    if genislik:
        im = im.resize((genislik, round(im.height * genislik / im.width)), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(b.getvalue()).decode()


def tr(x, fmt):
    if isinstance(x, (int, float)):
        bas = 0 if fmt == "#,##0" else 2
        return f"{x:,.{bas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "" if x is None else str(x)


def tablo_html(xlsx, satirlar, secili_sutun=None, secili_son=None):
    ws = load_workbook(xlsx).active
    hucreler = list(ws.iter_rows())
    harfler = "ABCDEFGH"[: ws.max_column]
    h = "<table><tr><th></th>" + "".join(f"<th>{x}</th>" for x in harfler) + "</tr>"
    onceki = 0
    for no in satirlar:
        if no - onceki > 1:
            h += "<tr class=atla><td class=n>…</td>" + "<td>…</td>" * len(harfler) + "</tr>"
        r = hucreler[no - 1]
        h += f"<tr><td class=n>{no}</td>" + "".join(
            f"<td class='{'r' if isinstance(c.value, (int, float)) else ''} {'b' if no == 1 else ''} "
            f"{'sec' if secili_sutun == j and 1 < no <= (secili_son or no) and isinstance(c.value, (int, float)) else ''}'>{tr(c.value, c.number_format)}</td>"
            for j, c in enumerate(r)) + "</tr>"
        onceki = no
    return h + "</table>", ws


STIL = """<link href="https://fonts.googleapis.com/css2?family=Inter+Tight:wght@700;800&family=Inter:wght@500;600&family=JetBrains+Mono:wght@500&display=swap" rel="stylesheet">
<style>
*{box-sizing:border-box;margin:0}
body{width:1280px;height:720px;overflow:hidden;background:radial-gradient(1200px 700px at 0% 0%,#1F2C42 0%,#101826 60%);
 color:#F4F1EA;font-family:Inter,sans-serif;padding:44px 52px}
h1{font:800 34px 'Inter Tight';letter-spacing:-.01em} h1 em{font-style:normal;color:#F2A93B}
.alt{color:#9AA7B8;font-size:17px;margin-top:6px}
.sahne{position:relative;margin-top:26px;height:560px}
.once{position:absolute;left:0;top:0;border-radius:8px;box-shadow:0 18px 36px rgba(0,0,0,.45)}
.ok{position:absolute;font:800 46px 'Inter Tight';color:#F2A93B}
.excel{position:absolute;right:0;top:0;background:#fff;color:#1d2733;border-radius:8px;overflow:hidden;box-shadow:0 22px 46px rgba(0,0,0,.5)}
.bar{background:#1E7145;color:#fff;font:600 13px Inter;padding:7px 12px}
.fx{display:flex;gap:10px;align-items:center;border-bottom:1px solid #DADDE1;padding:5px 10px;font-size:12px;color:#334}
.fx b{font:600 12px Inter;color:#667;border-right:1px solid #DADDE1;padding-right:10px}
.fx span{font-family:'JetBrains Mono'}
table{border-collapse:collapse;font-size:11.5px}
td,th{border:1px solid #DADDE1;padding:3px 6px;white-space:nowrap}
th{background:#F2F3F5;font-weight:600;color:#667;font-size:10.5px}
td.n{background:#F2F3F5;color:#778;text-align:center}
td.r{text-align:right;font-family:'JetBrains Mono';font-size:11px} td.b{font-weight:700;background:#F7FAF8}
td.sec{background:#E3F1E8;outline:1px solid #1E7145}
.durum{display:flex;justify-content:flex-end;gap:18px;background:#F2F3F5;border-top:1px solid #DADDE1;padding:5px 12px;font-size:12px;color:#334}
.damga{position:absolute;display:inline-flex;gap:8px;background:rgba(63,185,122,.14);border:1px solid rgba(63,185,122,.55);
 color:#7FE0AE;font-weight:600;font-size:15px;padding:7px 13px;border-radius:999px}
</style>"""


def fatura():
    xlsx = KOK / "calisma/portfolyo-fatura/cikti/portfolyo-fatura.xlsx"
    tb, ws = tablo_html(xlsx, list(range(1, 17)), secili_sutun=5, secili_son=13)
    toplam = sum(ws.cell(r, 6).value for r in range(2, 14))
    once = veri_uri(KOK / "spikes/s1-donusturme/veri/b_fatura_tarama_s1.jpg", (150, 150, 1030, 880), 400)
    return f"""<!doctype html><meta charset=utf-8>{STIL}<body>
<h1>Taranmış fatura → <em>Excel</em></h1><div class=alt>Rakamlar Excel'de sayı: sütunu seçince toplam doğrudan çıkıyor</div>
<div class=sahne><img class=once src="{once}" style="width:400px;top:20px"><div class=ok style="left:414px;top:200px">→</div>
<div class=excel style="width:700px"><div class=bar>fatura.xlsx</div><div class=fx><b>F2</b><span>{str(ws['F2'].value).replace('.', ',')}</span></div>{tb}
<div class=durum><span>Seçili: F2:F13</span><span>Toplam: {tr(toplam, '')}</span></div></div>
<span class=damga style="left:0;bottom:6px">✓ Satır hesapları ve KDV teslimden önce kontrol edildi</span></div></body>"""


def sayim():
    xlsx = KOK / "calisma/portfolyo-sayim/cikti/2026-10-23-depo-sayim.xlsx"
    tb, ws = tablo_html(xlsx, list(range(1, 10)) + [37, 38, 39, 40] + [58, 59, 60, 61])
    s = KOK / "calisma/portfolyo-sayim/sayfalar"
    s1 = veri_uri(s / "depo_sayim_taranmis_s1.png", (180, 150, 1500, 1880), 250)
    s2 = veri_uri(s / "depo_sayim_taranmis_s2.png", (180, 150, 1500, 1880), 250)
    return f"""<!doctype html><meta charset=utf-8>{STIL}<body>
<h1>İki sayfalık sayım tutanağı → <em>tek tablo</em></h1><div class=alt>Taranmış PDF, 60 satır; sayfa başlıkları ayıklandı, satırlar tek tabloda birleşti</div>
<div class=sahne><img class=once src="{s1}" style="width:250px;top:0"><img class=once src="{s2}" style="width:250px;left:150px;top:150px">
<div class=ok style="left:430px;top:220px">→</div>
<div class=excel style="width:640px"><div class=bar>depo-sayim.xlsx</div>{tb}<div class=durum><span>61 satır · tek çalışma sayfası</span></div></div>
<span class=damga style="right:0;bottom:6px">✓ 60/60 satırda miktar × fiyat = tutar</span></div></body>"""


with sync_playwright() as p:
    b = p.chromium.launch(executable_path=str(CHROME))
    for ad, html in [("portfolyo-fatura", fatura()), ("portfolyo-sayim", sayim())]:
        (OUT / f"{ad}.html").write_text(html, encoding="utf-8")
        s = b.new_page(viewport={"width": 1280, "height": 720})
        s.goto((OUT / f"{ad}.html").as_uri())
        s.wait_for_timeout(1500)
        s.screenshot(path=str(OUT / f"{ad}.png"))
        print(ad, "ok")
    b.close()
