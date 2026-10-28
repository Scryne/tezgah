"""Trendyol ürün düzenleme kiti: müşterinin Excel'i -> Claude düzenlemesi -> kurallı kontrol -> teslim.

Akış (sipariş klasörü `tezgah yeni` ile açılır, müşterinin Excel'i girdi/ altına konur):
1. `tezgah trendyol ac <klasor>`: Excel'i okur, urunler.json (mevcut hâl) + duzenleme_claude.json şablonu yazar.
   Excel yoksa (ücretsiz deneme) urunler.json'u Claude ürün sayfasından elle yazar.
2. Claude duzenleme_claude.json'u doldurur: başlık, açıklama, özellikler; her özelliğin kaynağı. Bilinmeyen bilgi `sorular`a.
3. `tezgah trendyol teslim <klasor>`: kontrol eder. Temizse teslim/ altına boyanmış Excel kopyası + Değişiklikler sayfası
   + kopyala-yapıştır metin yazar.

Kaynak kuralı ("bilmediğimi sorarım, uydurmam"): urun = değer mevcut ürün verisinde geçiyor; satici = satici.json'da aynı
değer var; gorsel = yalnız görselden okunabilen alanlar (renk, desen...). Tahmin kaynağı yoktur.
"""
import json
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .sayi import kucult

URUNLER = "urunler.json"
DUZENLEME = "duzenleme_claude.json"
SATICI = "satici.json"

BASLIK_ADLARI = ("ürün adı", "ürün ismi", "başlık", "urun adi")
ACIKLAMA_ADLARI = ("ürün açıklaması", "açıklama", "urun aciklamasi")
BARKOD_ADLARI = ("barkod", "barcode")
# Varsayım (doğrulanmadı): Trendyol ürün adı sınırı 100 karakter; müşterinin şablonu farklı derse buradan değişir.
BASLIK_EN_UZUN = 100
BASLIK_EN_KISA = 15
ACIKLAMA_EN_KISA = 150
GORSEL_ALANLAR = {"renk", "desen", "şekil", "tarz", "yaka tipi", "kol tipi", "boy"}
# Bizim kuralımız: pazaryeri metninde süreli/abartılı vaat ve iletişim bilgisi olmaz.
YASAK_IFADELER = ("en ucuz", "en iyi", "indirim", "kampanya", "bedava", "ücretsiz kargo", "stokta son")
_ILETISIM = re.compile(
    r"(\+?90[\s-]?)?0?5\d{2}[\s-]?\d{3}[\s-]?\d{2}[\s-]?\d{2}|[\w.+-]+@[\w-]+\.\w+|https?://|www\.|instagram|whatsapp",
    re.IGNORECASE,
)
BOYA = PatternFill("solid", start_color="FFE8F3D6")


@dataclass
class Sonuc:
    durum: str  # TEMIZ | INCELE | BEKLIYOR
    bulgular: list[str] = field(default_factory=list)
    sorular: list[str] = field(default_factory=list)
    ciktilar: list[Path] = field(default_factory=list)


def _excel(klasor: Path) -> Path | None:
    adaylar = sorted((klasor / "girdi").glob("*.xlsx"))
    return adaylar[0] if adaylar else None


def _baslik_satiri(ws) -> tuple[int, dict[str, int]]:
    """İlk 10 satırda barkod ya da ürün adı başlığı olan satırı bulur; {başlık: sütun no} döner."""
    for r in range(1, min(ws.max_row, 10) + 1):
        basliklar = {str(c.value).strip(): c.column for c in ws[r] if c.value is not None}
        adlar = {kucult(b) for b in basliklar}
        if adlar & set(BARKOD_ADLARI + BASLIK_ADLARI):
            return r, basliklar
    raise ValueError("Excel'de 'Barkod' ya da 'Ürün Adı' başlıklı satır bulunamadı (ilk 10 satır).")


def _sutun(basliklar: dict[str, int], adlar: tuple[str, ...]) -> str | None:
    return next((b for b in basliklar if kucult(b) in adlar), None)


def ac(klasor: Path) -> Path:
    """Excel'den urunler.json ve boş düzenleme şablonu yazar; Excel yoksa var olan urunler.json'u kullanır."""
    excel = _excel(klasor)
    if excel:
        ws = load_workbook(excel).active
        r0, basliklar = _baslik_satiri(ws)
        barkod = _sutun(basliklar, BARKOD_ADLARI)
        if not barkod:
            raise ValueError("Excel'de 'Barkod' sütunu yok; ürünler barkodla eşleştirilir.")
        urunler = []
        for r in range(r0 + 1, ws.max_row + 1):
            satir = {b: ws.cell(r, c).value for b, c in basliklar.items()}
            if satir.get(barkod) in (None, ""):
                continue
            urunler.append({"barkod": str(satir[barkod]), "alanlar": {k: v for k, v in satir.items() if v not in (None, "")}})
        (klasor / URUNLER).write_text(json.dumps(urunler, ensure_ascii=False, indent=1), encoding="utf-8")
    elif not (klasor / URUNLER).exists():
        raise ValueError(f"girdi/ altında Excel yok ve {URUNLER} yazılmamış.")
    urunler = json.loads((klasor / URUNLER).read_text(encoding="utf-8"))
    sablon = {"urunler": [{"barkod": u["barkod"], "baslik": "", "aciklama": "", "ozellikler": {}} for u in urunler],
              "sorular": []}
    hedef = klasor / DUZENLEME
    if not hedef.exists():
        hedef.write_text(json.dumps(sablon, ensure_ascii=False, indent=1), encoding="utf-8")
    return hedef


def _metin_kontrol(etiket: str, metin: str) -> list[str]:
    b = []
    if _ILETISIM.search(metin):
        b.append(f"{etiket}: iletişim bilgisi ya da bağlantı var")
    for ifade in YASAK_IFADELER:
        if ifade in kucult(metin):
            b.append(f"{etiket}: yasak ifade '{ifade}'")
    return b


def kontrol(klasor: Path) -> Sonuc:
    urunler = {u["barkod"]: u for u in json.loads((klasor / URUNLER).read_text(encoding="utf-8"))}
    d = json.loads((klasor / DUZENLEME).read_text(encoding="utf-8"))
    satici = json.loads((klasor / SATICI).read_text(encoding="utf-8")) if (klasor / SATICI).exists() else {}
    bulgular: list[str] = []
    gorulen = set()
    for u in d["urunler"]:
        k = u.get("barkod", "")
        if k not in urunler:
            bulgular.append(f"{k}: girdide olmayan barkod")
            continue
        gorulen.add(k)
        kaynak_metni = kucult(" ".join(str(v) for v in urunler[k]["alanlar"].values()))
        baslik, aciklama = u.get("baslik", "").strip(), u.get("aciklama", "").strip()
        if not BASLIK_EN_KISA <= len(baslik) <= BASLIK_EN_UZUN:
            bulgular.append(f"{k}: başlık {len(baslik)} karakter ({BASLIK_EN_KISA}–{BASLIK_EN_UZUN} olmalı)")
        harfler = [c for c in baslik if c.isalpha()]
        if harfler and sum(c.isupper() for c in harfler) / len(harfler) > 0.6:
            bulgular.append(f"{k}: başlık büyük harfle yazılmış")
        if len(aciklama) < ACIKLAMA_EN_KISA:
            bulgular.append(f"{k}: açıklama {len(aciklama)} karakter (en az {ACIKLAMA_EN_KISA})")
        bulgular += _metin_kontrol(f"{k} başlık", baslik) + _metin_kontrol(f"{k} açıklama", aciklama)
        for alan, o in u.get("ozellikler", {}).items():
            deger, kaynak = str(o.get("deger", "")).strip(), o.get("kaynak")
            if not deger:
                bulgular.append(f"{k} {alan}: boş değer")
            elif kaynak == "urun" and kucult(deger) not in kaynak_metni:
                bulgular.append(f"{k} {alan}: '{deger}' ürün verisinde geçmiyor (kaynak urun olamaz)")
            elif kaynak == "satici" and str(satici.get(k, {}).get(alan, "")).strip() != deger:
                bulgular.append(f"{k} {alan}: '{deger}' satıcı cevaplarında yok")
            elif kaynak == "gorsel" and kucult(alan) not in GORSEL_ALANLAR:
                bulgular.append(f"{k} {alan}: görselden okunamaz, satıcıya sorulmalı")
            elif kaynak not in ("urun", "satici", "gorsel"):
                bulgular.append(f"{k} {alan}: kaynak '{kaynak}' geçersiz (urun | satici | gorsel)")
    for k in urunler.keys() - gorulen:
        bulgular.append(f"{k}: düzenlemede yok")
    sorular = [s for s in d.get("sorular", []) if str(s).strip()]
    durum = "INCELE" if bulgular else "BEKLIYOR" if sorular else "TEMIZ"
    return Sonuc(durum, bulgular, sorular)


def teslim(klasor: Path) -> Sonuc:
    s = kontrol(klasor)
    if s.durum != "TEMIZ":
        return s
    d = json.loads((klasor / DUZENLEME).read_text(encoding="utf-8"))
    eski = {u["barkod"]: u["alanlar"] for u in json.loads((klasor / URUNLER).read_text(encoding="utf-8"))}
    cikti = klasor / "teslim"
    cikti.mkdir(exist_ok=True)

    satirlar = []
    for u in d["urunler"]:
        satirlar.append(f"## {u['baslik']}\n\nBarkod: {u['barkod']}\n\n**Ürün adı**\n\n{u['baslik']}\n\n**Açıklama**\n\n{u['aciklama']}\n")
        if u["ozellikler"]:
            satirlar.append("**Özellikler**\n\n" + "\n".join(f"- {a}: {o['deger']}" for a, o in u["ozellikler"].items()) + "\n")
    md = cikti / "urunler.md"
    md.write_text("\n".join(satirlar), encoding="utf-8")
    s.ciktilar.append(md)

    excel = _excel(klasor)
    if excel:
        hedef = cikti / f"{excel.stem}-duzenlendi.xlsx"
        shutil.copyfile(excel, hedef)
        wb = load_workbook(hedef)
        ws = wb.active
        r0, basliklar = _baslik_satiri(ws)
        barkod_c = basliklar[_sutun(basliklar, BARKOD_ADLARI)]
        satir_no = {str(ws.cell(r, barkod_c).value): r for r in range(r0 + 1, ws.max_row + 1)}
        degisen = wb.create_sheet("Değişiklikler")
        degisen.append(["Barkod", "Alan", "Eski", "Yeni"])
        elle = []
        for u in d["urunler"]:
            yeni = {_sutun(basliklar, BASLIK_ADLARI): u["baslik"], _sutun(basliklar, ACIKLAMA_ADLARI): u["aciklama"]}
            for alan, o in u["ozellikler"].items():
                sutun = _sutun(basliklar, (kucult(alan),))
                if sutun:
                    yeni[sutun] = o["deger"]
                else:
                    elle.append((u["barkod"], alan, o["deger"]))
            for sutun, deger in yeni.items():
                if sutun is None:
                    continue
                onceki = eski[u["barkod"]].get(sutun)
                if str(onceki or "") != deger:
                    h = ws.cell(satir_no[u["barkod"]], basliklar[sutun], deger)
                    h.fill = BOYA
                    degisen.append([u["barkod"], sutun, onceki, deger])
        for barkod, alan, deger in elle:
            degisen.append([barkod, f"{alan} (şablonda sütun yok, panelde elle)", None, deger])
        wb.save(hedef)
        s.ciktilar.append(hedef)
    return s


def soru_mesaji(sorular: list[str]) -> str:
    """Satıcıya gönderilecek Bionluk mesajı."""
    maddeler = "\n".join(f"{i}. {q}" for i, q in enumerate(sorular, 1))
    return f"Merhaba, ürünlerinizi düzenlerken uydurmamak için şu bilgileri sizden almam gerekiyor:\n{maddeler}\nCevaplarınızla hemen tamamlıyorum."
