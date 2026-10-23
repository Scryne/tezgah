"""Teslim dosyaları: Excel (sayılar gerçek sayı) ve Word."""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .sayi import coz, ondalik_basamak


def birlestir(tablolar: list[dict]) -> list[dict]:
    """Aynı belgede başlığı aynı olan ardışık tablolar (sayfaya bölünmüş tek tablo) tek tabloya iner;
    tekrarlanan başlık satırı düşer."""
    sonuc: list[dict] = []
    for t in tablolar:
        onceki = sonuc[-1] if sonuc else None
        if onceki and onceki["kaynak"] == t["kaynak"] and onceki["satirlar"][0] == t["satirlar"][0]:
            onceki["satirlar"].extend(t["satirlar"][1:])
            onceki["sayfalar"].append(t["sayfa"])
        else:
            sonuc.append({**t, "satirlar": [list(r) for r in t["satirlar"]], "sayfalar": [t["sayfa"]]})
    return sonuc


def _sayfa_adi(t: dict, kullanilan: set[str]) -> str:
    ad = re.sub(r"[\[\]:*?/\\]", "_", t["kaynak"])[:26] or "Tablo"
    aday, n = ad, 2
    while aday in kullanilan:
        aday, n = f"{ad[:24]}_{n}", n + 1
    kullanilan.add(aday)
    return aday


def excel_yaz(tablolar: list[dict], yol: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    kullanilan: set[str] = set()
    for t in tablolar:
        ws = wb.create_sheet(_sayfa_adi(t, kullanilan))
        for i, satir in enumerate(t["satirlar"], 1):
            for j, metin in enumerate(satir, 1):
                deger = coz(metin) if i > 1 else None
                hucre = ws.cell(i, j, deger if deger is not None else metin)
                if deger is not None:
                    hucre.number_format = "#,##0" if ondalik_basamak(metin) == 0 else "#,##0." + "0" * ondalik_basamak(metin)
                if i == 1:
                    hucre.font = Font(bold=True)
        for j in range(1, ws.max_column + 1):
            en_uzun = max((len(str(c.value or "")) for c in ws[get_column_letter(j)]), default=8)
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, en_uzun + 2), 60)
        ws.freeze_panes = "A2"
    wb.save(yol)


def word_yaz(tablolar: list[dict], yol: Path) -> None:
    doc = Document()
    for n, t in enumerate(tablolar):
        if n:
            doc.add_page_break()
        genislik = max(len(r) for r in t["satirlar"])
        tb = doc.add_table(rows=0, cols=genislik)
        tb.style = "Table Grid"
        for i, satir in enumerate(t["satirlar"]):
            hucreler = tb.add_row().cells
            for j in range(genislik):
                metin = satir[j] if j < len(satir) else ""
                run = hucreler[j].paragraphs[0].add_run(metin)
                run.font.size = Pt(9)
                run.bold = i == 0
    doc.save(yol)
